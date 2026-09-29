from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Vehicle-Speed-Checker-Beginner-Guide.docx"


def add_heading(doc, text, level=0):
    if level == 0:
        doc.add_heading(text, level=0)
    else:
        doc.add_heading(text, level=min(level, 3))


def add_bullet_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def add_section(doc, title, paragraphs=None, bullets=None):
    add_heading(doc, title, 1)
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if bullets:
        add_bullet_list(doc, bullets)


if __name__ == "__main__":
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Vehicle Speed Checker\nBeginner Guide")
    title_run.bold = True
    title_run.font.size = Pt(24)

    intro = [
        "This project is a small Python + Flask app that checks whether a vehicle speed is safe, at the limit, or over the speed limit.",
        "It is designed for learning: requirements analysis, domain logic, API design, testing, and GitHub-based workflow practice.",
    ]
    for paragraph in intro:
        doc.add_paragraph(paragraph)

    add_section(
        doc,
        "1. What this project does",
        paragraphs=[
            "The program compares a vehicle speed against a speed limit.",
            "It returns one of three results: SAFE, AT_LIMIT, or OVER_SPEED.",
            "It also rejects invalid inputs such as negative numbers, empty strings, None, booleans, or non-numeric values.",
        ],
        bullets=[
            "vehicle_speed < speed_limit -> SAFE",
            "vehicle_speed == speed_limit -> AT_LIMIT",
            "vehicle_speed > speed_limit -> OVER_SPEED",
        ],
    )

    add_section(
        doc,
        "2. Project structure",
        paragraphs=[
            "The project is split into source code and tests.",
            "The core business logic is in src/vehicle_speed_checker/checker.py.",
            "The API layer is in src/vehicle_speed_checker/api.py.",
        ],
        bullets=[
            "checker.py: validation and comparison logic",
            "api.py: Flask endpoint for POST /check",
            "tests/test_checker.py: direct logic tests",
            "tests/test_api.py: API and HTTP response tests",
            "README.md: setup and project overview",
        ],
    )

    add_section(
        doc,
        "3. Important files",
        paragraphs=[
            "The core function is check_compliance(vehicle_speed, speed_limit).",
            "The HTTP endpoint is POST /check.",
        ],
        bullets=[
            "src/vehicle_speed_checker/checker.py -> compares values and validates input",
            "src/vehicle_speed_checker/api.py -> converts HTTP requests to Python code",
            "docs/code-flow.md -> explains request flow step by step",
            "docs/swdd.md -> design description of the project",
        ],
    )

    add_section(
        doc,
        "4. How the app works",
        paragraphs=[
            "A client sends JSON with vehicle_speed and speed_limit.",
            "Flask reads the JSON body and passes the values into the checker.",
            "The checker validates both inputs, then compares them using Decimal values.",
        ],
    )

    add_code(
        doc,
        "Request example:\n{\n  \"vehicle_speed\": 79,\n  \"speed_limit\": 80\n}\n\nResponse:\n{\"status\": \"SAFE\"}",
    )

    add_section(
        doc,
        "5. Validation rules",
        paragraphs=[
            "Both inputs must be finite and non-negative numbers or numeric strings.",
            "The code rejects blank strings, None, booleans, negative numbers, and invalid text.",
        ],
        bullets=[
            "Accepted: 79, 80, 0, \"79\"",
            "Rejected: -1, None, True, \"\", \"abc\"",
        ],
    )

    add_section(
        doc,
        "6. How to run the project locally",
        paragraphs=[
            "Open PowerShell in the project folder.",
            "Create a virtual environment and install the package and test dependencies.",
        ],
    )
    add_code(
        doc,
        "cd C:\\Users\\qeu2kor\\Desktop\\info\\bosch-training\npython -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npython -m pip install --upgrade pip setuptools wheel\npython -m pip install --no-build-isolation -e \".[test]\"",
    )

    add_section(
        doc,
        "7. How to run tests",
        paragraphs=[
            "The project uses pytest.",
            "Run the test suite from the project root.",
        ],
    )
    add_code(doc, "python -m pytest")

    add_section(
        doc,
        "8. How to test the API manually",
        paragraphs=[
            "Start the Flask app in one terminal window.",
            "Then send a request from another terminal using PowerShell.",
        ],
    )
    add_code(
        doc,
        "flask --app vehicle_speed_checker.api run --debug\n\nInvoke-RestMethod -Method Post `\n  -Uri http://127.0.0.1:5000/check `\n  -ContentType \"application/json\" `\n  -Body '{\"vehicle_speed\":79,\"speed_limit\":80}'",
    )

    add_section(
        doc,
        "9. Beginner test strategy",
        paragraphs=[
            "Start with the simplest tests first.",
            "Check the happy path: 79 vs 80 gives SAFE.",
            "Then check boundaries: 80 equals limit and 81 is over speed.",
            "Finally test invalid inputs and malformed JSON.",
        ],
        bullets=[
            "Test valid values first",
            "Test boundary values second",
            "Test invalid values last",
            "Assert both status code and JSON body",
        ],
    )

    add_section(
        doc,
        "10. Summary",
        paragraphs=[
            "This project is a good beginner exercise because it combines Python, Flask, validation, business logic, and pytest testing in a small application.",
            "The main lesson is to keep the business logic separate from the web layer and test both clearly.",
        ],
    )

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
