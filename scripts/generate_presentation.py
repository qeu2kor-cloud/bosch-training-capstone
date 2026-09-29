from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Vehicle-Speed-Compliance-Checker.pptx"

INK = "18232B"
INK_SOFT = "26343D"
PAPER = "F4F2EC"
WHITE = "FFFEFA"
MUTED = "77838A"
LINE = "D8D8D0"
LIME = "C5E36B"
TEAL = "53BBA5"
CORAL = "EB755D"
PALE_GREEN = "E7EED2"
PALE_TEAL = "DCEEE8"
PALE_CORAL = "F5E3DC"

SW, SH = 13.333, 7.5
FONT = "Aptos"
DISPLAY = "Aptos Display"


def color(hex_value):
    return RGBColor.from_string(hex_value)


def box(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.06
    return shape


def label(slide, value, x, y, w, h, size=16, fill=INK, bold=False,
          font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
          margin=0, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(fill)
    return shape


def line(slide, x1, y1, x2, y2, fill=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(fill)
    shape.line.width = Pt(width)
    return shape


def dot(slide, x, y, diameter, fill, outline=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if outline:
        shape.line.color.rgb = color(outline)
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape


def base_slide(prs, section, title, subtitle=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = INK if dark else PAPER
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(bg)
    fg = WHITE if dark else INK
    accent = LIME if dark else TEAL
    label(slide, section.upper(), 0.72, 0.46, 7.8, 0.22, 10, accent, True)
    label(slide, title, 0.72, 0.88, 11.9, 0.58, 29, fg, True, DISPLAY)
    if subtitle:
        label(slide, subtitle, 0.74, 1.57, 11.8, 0.40, 13,
              "C5CED0" if dark else MUTED)
    line(slide, 0.72, 7.06, 12.61, 7.06,
         "3D4A50" if dark else LINE, 0.75)
    label(slide, "VEHICLE SPEED COMPLIANCE CHECKER", 0.74, 7.13,
          8.2, 0.16, 8, "AAB5B8" if dark else MUTED, True)
    label(slide, f"{len(prs.slides):02d}", 12.0, 7.11, 0.55, 0.19,
          9, accent, True, align=PP_ALIGN.RIGHT)
    return slide


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(INK)
    box(slide, 0, 0, 0.16, SH, LIME)
    label(slide, "SYSTEM BRIEFING     /     MVP 0.1.0", 0.82, 0.65,
          6.2, 0.25, 10, LIME, True)
    label(slide, "Vehicle Speed\nCompliance Checker", 0.82, 1.40,
          7.1, 1.65, 37, WHITE, True, DISPLAY)
    label(slide, "A clear, testable rule between a speed reading\nand its configured limit.",
          0.86, 3.38, 6.5, 0.75, 17, "C5CED0")
    line(slide, 0.86, 4.57, 6.8, 4.57, "526067", 1)
    label(slide, "PYTHON  /  FLASK  /  PYTEST", 0.86, 4.84,
          5.6, 0.28, 11, "C5CED0", True)
    label(slide, "DOMAIN LOGIC   +   HTTP API   +   CI", 0.86, 5.25,
          6.3, 0.28, 10, "92A0A4", True)

    # A simple, native-shape gauge makes the three-way comparison visible.
    for x, y, d, stroke in [(8.15, 1.08, 4.2, "34434B"),
                            (8.52, 1.45, 3.46, "46545A"),
                            (8.91, 1.84, 2.68, "5E6A6D")]:
        ring = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
        )
        ring.fill.background()
        ring.line.color.rgb = color(stroke)
        ring.line.width = Pt(1.4)
    dot(slide, 11.05, 2.70, 0.20, LIME)
    dot(slide, 9.12, 3.82, 0.16, TEAL)
    label(slide, "79", 9.22, 2.50, 2.1, 0.88, 52, WHITE, True,
          DISPLAY, align=PP_ALIGN.CENTER)
    label(slide, "km/h", 9.26, 3.38, 2.0, 0.28, 13,
          "AAB5B8", align=PP_ALIGN.CENTER)
    box(slide, 9.43, 4.12, 1.56, 0.48, LIME)
    label(slide, "SAFE", 9.43, 4.18, 1.56, 0.25, 12, INK, True,
          align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    label(slide, "79  <  80", 9.20, 5.02, 2.0, 0.32, 15,
          WHITE, True, align=PP_ALIGN.CENTER)
    line(slide, 0.86, 6.82, 12.55, 6.82, "3D4A50", 0.75)
    label(slide, "PROJECT OVERVIEW", 0.86, 6.96, 4.0, 0.18,
          8, "AAB5B8", True)
    label(slide, "01", 12.0, 6.94, 0.55, 0.2, 9, LIME, True,
          align=PP_ALIGN.RIGHT)


def add_story(prs):
    slide = base_slide(prs, "01 / The idea", "One comparison. Three clear outcomes.",
                       "A small service that turns two validated values into an unambiguous response.")
    columns = [
        ("01", "Receive", "A JSON request carries vehicle speed and the configured limit.", TEAL),
        ("02", "Validate", "Both values must be finite, numeric, and non-negative.", LIME),
        ("03", "Compare", "Decimal comparison maps the relationship to a named status.", CORAL),
    ]
    x_positions = [0.78, 4.76, 8.74]
    for (number, heading, body, accent), x in zip(columns, x_positions):
        label(slide, number, x, 2.45, 0.7, 0.42, 22, accent, True, DISPLAY)
        line(slide, x, 3.02, x + 3.18, 3.02, LINE, 1)
        label(slide, heading, x, 3.25, 3.1, 0.42, 22, INK, True, DISPLAY)
        label(slide, body, x, 3.86, 3.15, 0.88, 15, MUTED)
    box(slide, 0.78, 5.55, 11.7, 0.83, INK)
    label(slide, "DESIGN PRINCIPLE", 1.05, 5.78, 1.9, 0.20, 9, LIME, True)
    label(slide, "HTTP handles transport. The domain owns the rule.",
          3.04, 5.70, 8.8, 0.34, 18, WHITE, True, DISPLAY)


def add_architecture(prs):
    slide = base_slide(prs, "02 / Architecture", "A thin API around an independent domain",
                       "The route translates HTTP into a domain call, then maps the result back to JSON.")
    nodes = [
        (0.78, "CLIENT", "PowerShell\nor HTTP", TEAL),
        (3.86, "API ADAPTER", "api.py\nPOST /check", LIME),
        (6.94, "DOMAIN", "checker.py\nvalidate + compare", CORAL),
        (10.02, "RESPONSE", "JSON\n200 / 400", TEAL),
    ]
    for x, eyebrow, body, accent in nodes:
        box(slide, x, 2.83, 2.48, 1.54, WHITE, LINE)
        box(slide, x, 2.83, 0.08, 1.54, accent)
        label(slide, eyebrow, x + 0.24, 3.09, 2.05, 0.2, 9, MUTED, True)
        label(slide, body, x + 0.24, 3.43, 2.05, 0.60, 16, INK, True)
    for x in [3.31, 6.39, 9.47]:
        line(slide, x, 3.58, x + 0.43, 3.58, MUTED, 1.5)
        dot(slide, x + 0.37, 3.52, 0.12, MUTED)
    label(slide, "JSON request", 0.80, 4.62, 2.2, 0.24, 11, MUTED)
    label(slide, "check_compliance(speed, limit)", 4.02, 4.62,
          3.5, 0.24, 11, MUTED)
    label(slide, "status or validation error", 9.55, 4.62,
          2.5, 0.24, 11, MUTED)
    line(slide, 0.80, 5.28, 12.38, 5.28, LINE, 1)
    label(slide, "WHY IT MATTERS", 0.82, 5.60, 1.75, 0.22, 9, TEAL, True)
    label(slide, "Domain tests need no Flask server; API tests need no open network port.",
          2.64, 5.52, 9.2, 0.38, 16, INK, True)


def add_decision(prs):
    slide = base_slide(prs, "03 / Decision model", "The limit is the boundary",
                       "Inputs are converted to Decimal before comparison, avoiding binary float surprises.")
    rows = [
        ("79", "<", "80", "SAFE", PALE_TEAL, TEAL, "BELOW LIMIT"),
        ("80", "=", "80", "AT_LIMIT", PALE_GREEN, "739A35", "AT THE LIMIT"),
        ("81", ">", "80", "OVER_SPEED", PALE_CORAL, CORAL, "ABOVE LIMIT"),
    ]
    y_values = [2.43, 3.51, 4.59]
    for (speed, operator, limit, status, pale, accent, note), y in zip(rows, y_values):
        box(slide, 0.82, y, 11.65, 0.79, WHITE, LINE)
        box(slide, 0.82, y, 0.11, 0.79, accent)
        label(slide, speed, 1.20, y + 0.17, 1.10, 0.42, 25, INK, True, DISPLAY)
        label(slide, operator, 2.56, y + 0.17, 0.55, 0.42, 24, accent, True, DISPLAY)
        label(slide, limit, 3.35, y + 0.17, 1.10, 0.42, 25, INK, True, DISPLAY)
        label(slide, "VEHICLE SPEED", 1.20, y + 0.55, 1.65, 0.16, 8, MUTED, True)
        label(slide, "SPEED LIMIT", 3.35, y + 0.55, 1.55, 0.16, 8, MUTED, True)
        box(slide, 7.13, y + 0.18, 1.85, 0.42, pale)
        label(slide, status, 7.13, y + 0.23, 1.85, 0.25,
              11, accent, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        label(slide, note, 9.38, y + 0.26, 2.55, 0.22, 10, MUTED, True)
    label(slide, "Zero is valid: 0 vs 0 is AT_LIMIT; 0 vs 1 is SAFE.",
          0.86, 5.83, 8.2, 0.34, 15, INK, True)


def add_validation(prs):
    slide = base_slide(prs, "04 / Input contract", "Reject ambiguity before comparison",
                       "One normalization path validates both fields and returns a consistent error contract.")
    label(slide, "ACCEPT", 0.84, 2.33, 2.3, 0.28, 11, TEAL, True)
    label(slide, "REJECT", 6.84, 2.33, 2.3, 0.28, 11, CORAL, True)
    line(slide, 6.36, 2.34, 6.36, 5.75, LINE, 1)
    good = ["Finite numbers", "Numeric strings", "Zero and positive values"]
    bad = ["Missing or null values", "Negative values", "Booleans", "Empty / whitespace text",
           "Non-numeric or non-finite values"]
    for index, item in enumerate(good):
        y = 2.88 + index * 0.78
        dot(slide, 0.87, y + 0.05, 0.16, TEAL)
        label(slide, item, 1.22, y, 4.4, 0.34, 16, INK, True)
    for index, item in enumerate(bad):
        y = 2.83 + index * 0.60
        dot(slide, 6.86, y + 0.06, 0.15, CORAL)
        label(slide, item, 7.19, y, 4.8, 0.30, 14, INK, True)
    box(slide, 0.84, 5.77, 11.6, 0.60, INK)
    label(slide, "INVALID REQUEST", 1.08, 5.97, 1.8, 0.18, 9, LIME, True)
    label(slide, "HTTP 400", 3.08, 5.90, 1.7, 0.28, 15, WHITE, True)
    label(slide, '{"error": "vehicle_speed must be a non-negative number"}',
          5.00, 5.91, 6.8, 0.26, 12, "DCE3DF", font="Consolas")


def add_api(prs):
    slide = base_slide(prs, "05 / HTTP API", "Small surface. Predictable response.",
                       "POST /check accepts a JSON object and returns either a status or a field-level error.")
    box(slide, 0.82, 2.35, 5.58, 3.55, INK)
    label(slide, "REQUEST", 1.12, 2.67, 2.1, 0.2, 10, LIME, True)
    label(slide, "POST   /check", 1.12, 3.11, 3.8, 0.37, 19, WHITE, True, DISPLAY)
    line(slide, 1.12, 3.66, 5.98, 3.66, "47555B", 0.8)
    request = '{\n  "vehicle_speed": 79,\n  "speed_limit": 80\n}'
    label(slide, request, 1.12, 3.94, 4.7, 1.38,
          15, "E2E8E2", font="Consolas")
    label(slide, "Content-Type: application/json", 1.12, 5.45,
          4.8, 0.22, 10, "AAB5B8", font="Consolas")

    label(slide, "SUCCESS", 7.00, 2.52, 2.0, 0.20, 10, TEAL, True)
    label(slide, "200 OK", 7.00, 2.93, 3.0, 0.36, 23, INK, True, DISPLAY)
    box(slide, 7.00, 3.47, 4.85, 0.66, PALE_TEAL)
    label(slide, '{"status": "SAFE"}', 7.25, 3.65, 4.4, 0.25,
          15, INK, True, font="Consolas")
    line(slide, 7.00, 4.49, 12.02, 4.49, LINE, 1)
    label(slide, "INVALID INPUT", 7.00, 4.77, 2.4, 0.20, 10, CORAL, True)
    label(slide, "400 Bad Request", 7.00, 5.13, 3.5, 0.35,
          21, INK, True, DISPLAY)
    label(slide, "Malformed or non-object JSON is rejected before the domain call.",
          7.00, 5.62, 5.0, 0.55, 13, MUTED)


def add_quality(prs):
    slide = base_slide(prs, "06 / Verification", "Confidence at two useful levels",
                       "Fast domain checks protect the rule; test-client checks protect the HTTP contract.")
    box(slide, 0.82, 2.45, 5.60, 3.48, WHITE, LINE)
    box(slide, 0.82, 2.45, 5.60, 0.10, TEAL)
    label(slide, "DOMAIN", 1.15, 2.82, 2.3, 0.20, 10, TEAL, True)
    label(slide, "test_checker.py", 1.15, 3.22, 4.6, 0.38,
          21, INK, True, DISPLAY)
    label(slide, "Direct calls to check_compliance", 1.15, 3.78,
          4.6, 0.32, 14, MUTED)
    label(slide, "Boundaries 79 / 80 / 81", 1.15, 4.43,
          4.6, 0.28, 13, INK, True)
    label(slide, "Zero, numeric strings, invalid partitions", 1.15,
          4.88, 4.75, 0.48, 13, INK, True)

    box(slide, 6.82, 2.45, 5.60, 3.48, WHITE, LINE)
    box(slide, 6.82, 2.45, 5.60, 0.10, CORAL)
    label(slide, "HTTP CONTRACT", 7.15, 2.82, 2.7, 0.20, 10, CORAL, True)
    label(slide, "test_api.py", 7.15, 3.22, 4.6, 0.38,
          21, INK, True, DISPLAY)
    label(slide, "Flask application test client", 7.15, 3.78,
          4.6, 0.32, 14, MUTED)
    label(slide, "Status codes + complete JSON", 7.15, 4.43,
          4.6, 0.28, 13, INK, True)
    label(slide, "No server, socket, or external service", 7.15,
          4.88, 4.8, 0.48, 13, INK, True)
    box(slide, 0.82, 6.15, 11.60, 0.48, INK)
    label(slide, "CI", 1.08, 6.29, 0.55, 0.17, 10, LIME, True)
    label(slide, "pytest runs in GitHub Actions across Python 3.9–3.12.",
          1.85, 6.25, 9.8, 0.24, 13, WHITE, True)


def add_delivery(prs):
    slide = base_slide(prs, "07 / Delivery", "An intentionally focused MVP", dark=True,
                       subtitle="The repository is built for requirements, testing, development, review, and CI practice.")
    label(slide, "IN THE MVP", 0.86, 2.38, 2.2, 0.24, 10, LIME, True)
    label(slide, "Python domain function", 0.86, 2.89, 4.7, 0.32, 17, WHITE, True)
    label(slide, "Flask POST /check API", 0.86, 3.42, 4.7, 0.32, 17, WHITE, True)
    label(slide, "Parameterized pytest coverage", 0.86, 3.95, 4.7, 0.32, 17, WHITE, True)
    label(slide, "GitHub Actions CI", 0.86, 4.48, 4.7, 0.32, 17, WHITE, True)
    line(slide, 6.60, 2.38, 6.60, 5.18, "46545A", 1)
    label(slide, "OUT OF SCOPE", 7.05, 2.38, 2.5, 0.24, 10, CORAL, True)
    label(slide, "Authentication", 7.05, 2.91, 4.5, 0.30, 16, "D1D9D8")
    label(slide, "Persistence and user management", 7.05, 3.42, 5.0, 0.30, 16, "D1D9D8")
    label(slide, "Monitoring and rate limiting", 7.05, 3.93, 5.0, 0.30, 16, "D1D9D8")
    label(slide, "AWS infrastructure / deployment", 7.05, 4.44, 5.0, 0.30, 16, "D1D9D8")
    box(slide, 0.86, 5.60, 11.62, 0.72, INK_SOFT)
    label(slide, "NEXT STEP", 1.12, 5.86, 1.25, 0.17, 9, LIME, True)
    label(slide, "Harden and deploy only when the exercise calls for production concerns.",
          2.62, 5.78, 9.2, 0.32, 15, WHITE, True)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "Vehicle Speed Compliance Checker"
    prs.core_properties.subject = "Project overview, API behavior, architecture, and test strategy"
    prs.core_properties.author = "Vehicle Speed Checker Team"
    prs.core_properties.keywords = "Python, Flask, pytest, vehicle speed, API"

    add_cover(prs)
    add_story(prs)
    add_architecture(prs)
    add_decision(prs)
    add_validation(prs)
    add_api(prs)
    add_quality(prs)
    add_delivery(prs)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()